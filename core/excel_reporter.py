"""Excel 报告生成器模块

将解析结果（ParseResult 列表）输出为格式化的 Excel 报表，
包含表头样式、列宽、冻结首行、自动筛选等格式化功能。
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import ParseResult


class ExcelReporter:
    """Excel 报告生成器

    将多个文件的标注解析结果汇总为一份格式化的 Excel 报表。
    """

    # 列定义：(列标题, 列宽)
    COLUMNS: list[tuple[str, int]] = [
        ("序号", 8),
        ("标注内容", 35),
        ("X坐标", 15),
        ("Y坐标", 15),
        ("标注类型", 15),
        ("图层名", 20),
        ("来源文件", 30),
    ]

    # 表头字体：微软雅黑 11号 加粗 白色
    HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")

    # 表头填充：蓝色背景 #4472C4
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # 表头对齐：居中 自动换行
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 数据单元格字体：微软雅黑 10号
    CELL_FONT = Font(name="微软雅黑", size=10)

    # 数据单元格对齐：垂直居中 自动换行
    CELL_ALIGNMENT = Alignment(vertical="center", wrap_text=True)

    # 四边细线边框
    THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def generate_report(
        self,
        results: list[ParseResult],
        output_path: str,
        progress_callback=None,
    ) -> str:
        """生成 Excel 汇总报告

        Args:
            results: 多个文件的解析结果列表
            output_path: 输出文件路径
            progress_callback: 进度回调函数，签名为 callback(current, total)

        Returns:
            输出文件的绝对路径字符串
        """
        # 1. 收集所有 ParseResult 中的 items 为扁平列表
        all_items = []
        for result in results:
            if result.error is not None:
                continue
            all_items.extend(result.items)

        # 2. 重新全局连续编号
        for i, item in enumerate(all_items, start=1):
            item.index = i

        total = len(all_items)

        # 3. 构建 pandas DataFrame
        records = []
        for i, item in enumerate(all_items):
            records.append({
                "序号": item.index,
                "标注内容": item.content,
                "X坐标": round(item.x, 4),
                "Y坐标": round(item.y, 4),
                "标注类型": item.annotation_type.value,
                "图层名": item.layer,
                "来源文件": item.source_file,
            })

            # 进度回调：收集阶段
            if progress_callback is not None:
                progress_callback(i + 1, total)

        df = pd.DataFrame(records, columns=[col[0] for col in self.COLUMNS])

        # 4. 确保输出目录存在
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        # 5. 写入 Excel（使用 openpyxl 引擎）
        df.to_excel(str(output), index=False, engine="openpyxl")

        # 6. 应用格式化
        self._apply_formatting(str(output))

        return str(output.resolve())

    def _apply_formatting(self, file_path: str) -> None:
        """对已生成的 Excel 文件应用格式化样式

        Args:
            file_path: Excel 文件路径
        """
        wb = load_workbook(file_path)
        ws = wb.active

        # 2. 表头行设置：字体、填充、对齐、边框
        for col_idx in range(1, len(self.COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.THIN_BORDER

        # 3. 数据行设置：字体、对齐、边框
        max_row = ws.max_row
        for row_idx in range(2, max_row + 1):
            for col_idx in range(1, len(self.COLUMNS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = self.CELL_FONT
                cell.alignment = self.CELL_ALIGNMENT
                cell.border = self.THIN_BORDER

        # 4. 列宽按 COLUMNS 定义设置
        for col_idx, (_, width) in enumerate(self.COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # 5. 冻结首行 A2
        ws.freeze_panes = "A2"

        # 6. 添加自动筛选
        ws.auto_filter.ref = ws.dimensions

        # 7. 保存并关闭
        wb.save(file_path)
        wb.close()
